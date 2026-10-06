import re
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook


def main() -> int:
    workbook_path = Path(sys.argv[1])
    sheet = load_workbook(workbook_path, data_only=False).active
    rows = list(sheet.iter_rows(min_row=8, max_row=85, values_only=True))

    tarot_names = [row[0] for row in rows]
    mapped_cards = [row[1] for row in rows]
    duplicate_tarot = [name for name, count in Counter(tarot_names).items() if count > 1]
    duplicate_mapped = [name for name, count in Counter(mapped_cards).items() if count > 1]

    missing_files = []
    invalid_links = []
    for row in rows:
        formula = row[4] or ""
        match = re.search(r'file:///([^\"]+)', formula)
        if not match:
            invalid_links.append(row[0])
            continue
        art_path = Path(unquote(match.group(1)).replace("/", "\\"))
        if not art_path.exists():
            missing_files.append((row[0], str(art_path)))

    print(f"rows: {len(rows)}")
    print(f"unique tarot archetypes: {len(set(tarot_names))}")
    print(f"unique mapped cards: {len(set(mapped_cards))}")
    print(f"duplicate tarot archetypes: {duplicate_tarot}")
    print(f"duplicate mapped cards: {duplicate_mapped}")
    print(f"sub-areas: {dict(Counter(row[2] for row in rows))}")
    print(f"types: {dict(Counter(row[3] for row in rows))}")
    print(f"art status: {dict(Counter(row[5] for row in rows))}")
    print(f"mapping state: {dict(Counter(row[6] for row in rows))}")
    print(f"source status: {dict(Counter(row[7] for row in rows))}")
    print(f"invalid art links: {len(invalid_links)}")
    print(f"missing art files: {len(missing_files)}")
    for tarot_name, path in missing_files:
        print(f"MISSING: {tarot_name}: {path}")

    return int(
        len(rows) != 78
        or len(set(tarot_names)) != 78
        or len(set(mapped_cards)) != 78
        or bool(invalid_links)
        or bool(missing_files)
    )


if __name__ == "__main__":
    raise SystemExit(main())
