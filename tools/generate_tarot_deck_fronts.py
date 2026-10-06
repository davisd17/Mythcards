import re
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from generate_tarot_front_proof import render_card


ROOT = Path(__file__).resolve().parents[1]
TRACKER = (
    ROOT
    / "outputs"
    / "019f2e9a-7504-7ee0-9b75-0d4456667415"
    / "russian_tarot_deck_tracker.xlsx"
)
OUTPUT_DIR = Path(r"D:\MythCards-working-images\card-fronts\russian-tarot-deck")

MAJOR_ARCANA = [
    "The Fool",
    "The Magician",
    "The High Priestess",
    "The Empress",
    "The Emperor",
    "The Hierophant",
    "The Lovers",
    "The Chariot",
    "Strength",
    "The Hermit",
    "Wheel Of Fortune",
    "Justice",
    "The Hanged Man",
    "Death",
    "Temperance",
    "The Devil",
    "The Tower",
    "The Star",
    "The Moon",
    "The Sun",
    "Judgement",
    "The World",
]
ROMAN = [
    "0",
    "I",
    "II",
    "III",
    "IV",
    "V",
    "VI",
    "VII",
    "VIII",
    "IX",
    "X",
    "XI",
    "XII",
    "XIII",
    "XIV",
    "XV",
    "XVI",
    "XVII",
    "XVIII",
    "XIX",
    "XX",
    "XXI",
]


def art_path_from_formula(formula: str) -> Path:
    match = re.search(r'file:///([^\"]+)', formula)
    if not match:
        raise ValueError(f"Invalid art hyperlink: {formula}")
    return Path(unquote(match.group(1)).replace("/", "\\"))


def display_title(tarot_name: str) -> str:
    if tarot_name in MAJOR_ARCANA:
        index = MAJOR_ARCANA.index(tarot_name)
        return f"{ROMAN[index]}  -  {tarot_name.upper()}"
    return tarot_name.upper()


def filename(index: int, tarot_name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", tarot_name.lower()).strip("-")
    return f"{index:02d}-{slug}.png"


def make_contact_sheet(paths: list[Path]) -> Path:
    columns = 6
    thumb_size = (210, 360)
    gap = 16
    rows = (len(paths) + columns - 1) // columns
    sheet = Image.new(
        "RGB",
        (columns * thumb_size[0] + (columns + 1) * gap, rows * thumb_size[1] + (rows + 1) * gap),
        "#25282A",
    )
    for index, path in enumerate(paths):
        image = Image.open(path).convert("RGB")
        image.thumbnail(thumb_size, Image.Resampling.LANCZOS)
        x = gap + (index % columns) * (thumb_size[0] + gap)
        y = gap + (index // columns) * (thumb_size[1] + gap)
        sheet.paste(image, (x, y))
    output = OUTPUT_DIR / "russian-tarot-fronts-contact-sheet.jpg"
    sheet.save(output, quality=90, optimize=True)
    return output


def main() -> None:
    sheet = load_workbook(TRACKER, data_only=False).active
    rows = list(sheet.iter_rows(min_row=8, max_row=85, values_only=True))
    if len(rows) != 78:
        raise ValueError(f"Expected 78 tracker rows, found {len(rows)}")

    outputs = []
    for index, row in enumerate(rows):
        tarot_name = row[0]
        source = art_path_from_formula(row[4])
        output = OUTPUT_DIR / filename(index, tarot_name)
        render_card(source, output, display_title(tarot_name))
        outputs.append(output)
        print(f"{index + 1:02d}/78 {tarot_name}: {output.name}")

    contact_sheet = make_contact_sheet(outputs)
    print(f"Contact sheet: {contact_sheet}")


if __name__ == "__main__":
    main()
