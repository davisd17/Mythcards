import json
import re
from pathlib import Path
from urllib.parse import unquote, urlparse

from PIL import Image, ImageDraw, ImageFont, ImageOps
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TRACKER = (
    ROOT
    / "outputs"
    / "019f2e9a-7504-7ee0-9b75-0d4456667415"
    / "russian_tarot_deck_tracker.xlsx"
)
OUTPUT = ROOT / "artifacts" / "russian-tarot-deck-contact-sheet.png"

COLS = 13
ROWS = 6
TILE_W = 350
TILE_H = 550
ART_H = 468
GUTTER = 18
MARGIN = 48
HEADER_H = 150


def art_path(value: str) -> Path:
    match = re.search(r"linkLocation=(file:///.*?), friendlyName=", value)
    if not match:
        raise ValueError(f"Cannot parse art hyperlink: {value}")
    parsed = urlparse(match.group(1))
    return Path(unquote(parsed.path.lstrip("/")))


def tracker_rows() -> list[dict]:
    workbook = load_workbook(TRACKER, data_only=True, read_only=False)
    sheet = workbook["Russian tarot"]
    rows = []
    for row in range(8, 86):
        tarot_name = sheet.cell(row, 1).value
        if not tarot_name:
            continue
        rows.append(
            {
                "tarot": tarot_name,
                "mythcards_card": sheet.cell(row, 2).value,
                "sub_area": sheet.cell(row, 3).value,
                "card_type": sheet.cell(row, 4).value,
                "approved_asset": art_path(sheet.cell(row, 5).value),
            }
        )
    return rows


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / name), size)


def fit_text(draw: ImageDraw.ImageDraw, text: str, max_width: int, start: int, bold: bool = False):
    size = start
    while size > 12:
        font = load_font(size, bold)
        if draw.textbbox((0, 0), text, font=font)[2] <= max_width:
            return font
        size -= 1
    return load_font(12, bold)


def render_placeholder(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], label: str) -> None:
    left, top, right, bottom = box
    draw.rectangle(box, fill="#20252A")
    draw.line((left + 24, top + 24, right - 24, bottom - 24), fill="#414950", width=3)
    draw.line((right - 24, top + 24, left + 24, bottom - 24), fill="#414950", width=3)
    font = load_font(23, True)
    bbox = draw.textbbox((0, 0), label, font=font)
    draw.rectangle(
        (left + 18, (top + bottom) // 2 - 28, right - 18, (top + bottom) // 2 + 28),
        fill="#20252A",
    )
    draw.text(
        ((left + right - (bbox[2] - bbox[0])) // 2, (top + bottom - (bbox[3] - bbox[1])) // 2),
        label,
        font=font,
        fill="#C7CDD1",
    )


def main() -> None:
    mappings = tracker_rows()
    if len(mappings) != 78:
        raise ValueError(f"Expected 78 tracker rows, found {len(mappings)}")

    width = MARGIN * 2 + COLS * TILE_W + (COLS - 1) * GUTTER
    height = HEADER_H + MARGIN + ROWS * TILE_H + (ROWS - 1) * GUTTER + MARGIN
    canvas = Image.new("RGB", (width, height), "#111416")
    draw = ImageDraw.Draw(canvas)

    title_font = load_font(54, True)
    subtitle_font = load_font(25)
    draw.text((MARGIN, 34), "MYTHCARDS RUSSIAN TAROT", font=title_font, fill="#F1EEE6")
    draw.text(
        (MARGIN, 98),
        "78-card visual tracker | preferred artwork where linked | canonical tarot order",
        font=subtitle_font,
        fill="#AAB2B7",
    )

    for index, entry in enumerate(mappings):
        tarot_name = entry["tarot"]
        row, col = divmod(index, COLS)
        x = MARGIN + col * (TILE_W + GUTTER)
        y = HEADER_H + MARGIN + row * (TILE_H + GUTTER)
        art_box = (x, y, x + TILE_W, y + ART_H)
        asset = entry["approved_asset"]

        if asset and asset.exists():
            with Image.open(asset) as source:
                art = ImageOps.fit(
                    source.convert("RGB"),
                    (TILE_W, ART_H),
                    method=Image.Resampling.LANCZOS,
                    centering=(0.5, 0.42),
                )
            canvas.paste(art, (x, y))
        else:
            render_placeholder(draw, art_box, "MISSING ART")

        draw.rectangle((x, y, x + TILE_W - 1, y + TILE_H - 1), outline="#6F767A", width=2)
        draw.rectangle((x, y + ART_H, x + TILE_W, y + TILE_H), fill="#181C1F")

        number = f"{index + 1:02d}"
        name_font = fit_text(draw, tarot_name.upper(), TILE_W - 32, 22, True)
        draw.text((x + 14, y + ART_H + 10), number, font=load_font(17, True), fill="#8F989E")
        draw.text((x + 14, y + ART_H + 34), tarot_name.upper(), font=name_font, fill="#F0ECE3")

        subject_text = entry["mythcards_card"]
        subject_font = fit_text(draw, subject_text, TILE_W - 28, 15)
        draw.text((x + 14, y + TILE_H - 22), subject_text, font=subject_font, fill="#AEB6BB", anchor="ls")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(OUTPUT, optimize=True)
    print(OUTPUT)


if __name__ == "__main__":
    main()
